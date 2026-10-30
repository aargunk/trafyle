"""Dışa aktarma: düz metin, zaman damgalı metin, SRT, VTT, Word ve PDF.

Konuşmacı bilgisi varsa her biçimde konuşmacı adı da yazılır.
"""
from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Sequence

from engine import Segment


# ---------------------------------------------------------------- zaman biçimleri

def fmt_clock(sec: float) -> str:
    """01:02:03"""
    sec = max(0, int(round(sec)))
    return f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


def fmt_editor(sec: float) -> str:
    """Düzeltme ekranındaki damga: 01:02:03.4"""
    sec = max(0.0, sec)
    tenths = int(round(sec * 10))
    s, t = divmod(tenths, 10)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}.{t}"


def _fmt_sub(sec: float, sep: str) -> str:
    ms = max(0, int(round(sec * 1000)))
    s, ms = divmod(ms, 1000)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}{sep}{ms:03d}"


# ---------------------------------------------------------------- düzeltme ekranı metni

LINE_RE = re.compile(r"^\[(\d{1,2}):(\d{2}):(\d{2})(?:[.,](\d{1,3}))?\]\s?(.*)$")


def editor_line(seg: Segment) -> tuple[str, str, str]:
    """(damga, konuşmacı öneki, metin) — düzeltme ekranında ayrı renklendirmek için."""
    return f"[{fmt_editor(seg.start)}]", (f"{seg.speaker}: " if seg.speaker else ""), seg.text


def to_editor_text(segments: Sequence[Segment]) -> str:
    return "\n".join("{} {}{}".format(*editor_line(s)) for s in segments)


def split_speaker(body: str, speakers: Sequence[str]) -> tuple[str, str]:
    """Satırın başındaki "Ad: " önekini ayırır. Yalnızca bilinen konuşmacı adları tanınır,
    böylece "Şöyle ki: …" gibi cümleler konuşmacı sanılmaz."""
    for name in sorted(speakers, key=len, reverse=True):
        if name and body.startswith(name + ":"):
            return name, body[len(name) + 1:].strip()
    return "", body


def parse_editor_text(text: str, duration: float, original: Sequence[Segment] = (),
                      speakers: Sequence[str] = ()) -> list[Segment]:
    """Düzeltme ekranındaki metni tekrar parçalara çevirir.

    Her satır "[SS:DD:ss.s] Konuşmacı: metin" biçimindedir (konuşmacı isteğe bağlı).
    Kullanıcı damgasız bir satır eklerse önceki parçaya eklenir. Damgası değişmeyen satırlar
    özgün bitiş zamanını korur; diğerlerinde bitiş, sonraki satırın başıdır.
    """
    orig_end = {round(s.start, 1): s.end for s in original}
    items: list[list] = []  # [start, speaker, text]
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = LINE_RE.match(line)
        if m:
            h, mi, s, frac, body = m.groups()
            start = int(h) * 3600 + int(mi) * 60 + int(s) + (int(frac) / (10 ** len(frac)) if frac else 0)
            spk, body = split_speaker(body.strip(), speakers)
            items.append([start, spk, body.strip()])
        elif items:
            items[-1][2] = (items[-1][2] + " " + line.strip()).strip()
        else:
            items.append([0.0, "", line.strip()])
    items = [it for it in items if it[2]]
    items.sort(key=lambda it: it[0])
    segs: list[Segment] = []
    for i, (start, spk, body) in enumerate(items):
        nxt = items[i + 1][0] if i + 1 < len(items) else (duration if duration > start else start + 2.0)
        end = min(orig_end.get(round(start, 1), nxt), nxt)
        if end <= start:
            end = start + 0.5
        segs.append(Segment(start, end, body, spk))
    return segs


# ---------------------------------------------------------------- paragraflar

def paragraphs(segments: Sequence[Segment], gap: float = 2.0, max_chars: int = 700) -> list[tuple[float, str, str]]:
    """Parçaları okunur paragraflara toplar: (başlangıç, konuşmacı, metin).

    Konuşmacı değişince, uzun sessizlikte ya da paragraf çok uzayınca yeni paragraf başlar.
    """
    out: list[tuple[float, str, str]] = []
    cur: list[str] = []
    start = 0.0
    spk = ""
    prev_end = None
    for s in segments:
        new_speaker = cur and s.speaker != spk
        long_gap = prev_end is not None and s.start - prev_end >= gap
        too_long = sum(len(c) + 1 for c in cur) >= max_chars and cur and cur[-1].endswith((".", "?", "!", "…"))
        if cur and (new_speaker or long_gap or too_long):
            out.append((start, spk, " ".join(cur)))
            cur = []
        if not cur:
            start, spk = s.start, s.speaker
        cur.append(s.text)
        prev_end = s.end
    if cur:
        out.append((start, spk, " ".join(cur)))
    return out


def _with_speaker(speaker: str, text: str) -> str:
    return f"{speaker}: {text}" if speaker else text


# ---------------------------------------------------------------- metin biçimleri

def to_txt(segments: Sequence[Segment], timestamps: bool = False) -> str:
    if timestamps:
        return "\n".join(f"[{fmt_clock(s.start)}] {_with_speaker(s.speaker, s.text)}" for s in segments) + "\n"
    return "\n\n".join(_with_speaker(spk, p) for _, spk, p in paragraphs(segments)) + "\n"


def to_srt(segments: Sequence[Segment]) -> str:
    blocks = []
    for i, s in enumerate(segments, 1):
        blocks.append(f"{i}\n{_fmt_sub(s.start, ',')} --> {_fmt_sub(s.end, ',')}\n{_with_speaker(s.speaker, s.text)}\n")
    return "\n".join(blocks)


def to_vtt(segments: Sequence[Segment]) -> str:
    body = "\n".join(
        f"{_fmt_sub(s.start, '.')} --> {_fmt_sub(s.end, '.')}\n"
        + (f"<v {s.speaker}>{s.text}" if s.speaker else s.text) + "\n"
        for s in segments)
    return "WEBVTT\n\n" + body


def _meta_line(meta: dict) -> str:
    parts = []
    if meta.get("duration"):
        parts.append(f"Süre: {fmt_clock(meta['duration'])}")
    if meta.get("speakers"):
        parts.append(f"Konuşmacılar: {', '.join(meta['speakers'])}")
    if meta.get("model"):
        parts.append(f"Model: {meta['model']}")
    parts.append(f"Tarih: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    return " · ".join(parts)


# ---------------------------------------------------------------- Word

def to_docx(path: str, segments: Sequence[Segment], title: str, meta: dict, timestamps: bool = False) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(title, level=1)
    p = doc.add_paragraph(_meta_line(meta))
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    def add(ts: float | None, speaker: str, text: str):
        para = doc.add_paragraph()
        if ts is not None:
            r = para.add_run(f"[{fmt_clock(ts)}] ")
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            r.font.size = Pt(9)
        if speaker:
            r = para.add_run(f"{speaker}: ")
            r.bold = True
        para.add_run(text)

    if timestamps:
        for s in segments:
            add(s.start, s.speaker, s.text)
    else:
        for _, spk, text in paragraphs(segments):
            add(None, spk, text)
    doc.save(path)


# ---------------------------------------------------------------- PDF

_FONT_CANDIDATES = [
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/Library/Fonts/Arial Unicode.ttf", None),
    ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", None),
    ("/System/Library/Fonts/Supplemental/Verdana.ttf", "/System/Library/Fonts/Supplemental/Verdana Bold.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]


def _find_font() -> tuple[str, str | None] | None:
    for regular, bold in _FONT_CANDIDATES:
        if os.path.isfile(regular):
            return regular, (bold if bold and os.path.isfile(bold) else None)
    return None


def to_pdf(path: str, segments: Sequence[Segment], title: str, meta: dict, timestamps: bool = False) -> None:
    from fpdf import FPDF

    font = _find_font()
    if not font:
        raise RuntimeError("PDF için Türkçe karakterleri destekleyen bir yazı tipi bulunamadı. Word olarak dışa aktarın.")
    regular, bold = font
    B = "B" if bold else ""

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(True, margin=20)
    pdf.add_font("Govde", "", regular)
    if bold:
        pdf.add_font("Govde", "B", bold)
    pdf.add_page()

    pdf.set_font("Govde", B, 16)
    pdf.multi_cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Govde", "", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.multi_cell(0, 5, _meta_line(meta), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_text_color(0, 0, 0)

    def add(ts: float | None, speaker: str, text: str, after: float):
        if ts is not None:
            pdf.set_text_color(100, 100, 100)
            pdf.set_font("Govde", "", 9)
            pdf.cell(22, 6, f"[{fmt_clock(ts)}]")
            pdf.set_text_color(0, 0, 0)
        if speaker:
            # Konuşmacı adı kalın, metin normal; markdown ile tek hücrede yazılır.
            safe = text.replace("**", "*\u200b*").replace("__", "_\u200b_")
            pdf.set_font("Govde", "", 11)
            pdf.multi_cell(0, 6, f"**{speaker}:** {safe}" if bold else f"{speaker}: {text}",
                           markdown=bool(bold), new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Govde", "", 11)
            pdf.multi_cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(after)

    if timestamps:
        for s in segments:
            add(s.start, s.speaker, s.text, 1)
    else:
        for _, spk, text in paragraphs(segments):
            add(None, spk, text, 3)
    pdf.output(path)


# ---------------------------------------------------------------- tek giriş noktası

FORMATS = [
    # (anahtar, menü adı, uzantı)
    ("docx", "Word (.docx)", ".docx"),
    ("pdf", "PDF (.pdf)", ".pdf"),
    ("txt", "Düz metin (.txt)", ".txt"),
    ("srt", "Altyazı (.srt)", ".srt"),
    ("vtt", "Altyazı (.vtt)", ".vtt"),
]


def export(fmt: str, path: str, segments: Sequence[Segment], title: str, meta: dict, timestamps: bool) -> None:
    if fmt == "docx":
        to_docx(path, segments, title, meta, timestamps)
        return
    if fmt == "pdf":
        to_pdf(path, segments, title, meta, timestamps)
        return
    if fmt == "txt":
        content = to_txt(segments, timestamps)
    elif fmt == "srt":
        content = to_srt(segments)
    elif fmt == "vtt":
        content = to_vtt(segments)
    else:
        raise ValueError(fmt)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


# ------------------------------------------------------------------ özet (sade Markdown) çıktıları

_MD_INLINE = re.compile(r"\*\*(.+?)\*\*")


def _md_blocks(md: str) -> list[tuple[str, str]]:
    """Sade Markdown'u (tür, metin) bloklarına ayırır: h1/h2/h3, li, p."""
    out = []
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        s = line.lstrip()
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            out.append(("h%d" % min(3, len(m.group(1))), m.group(2).strip()))
        elif re.match(r"^([-*•]|\d+[.)])\s+", s):
            out.append(("li", re.sub(r"^([-*•]|\d+[.)])\s+", "", s)))
        else:
            out.append(("p", s))
    return out


def md_to_text(md: str) -> str:
    lines = []
    for kind, text in _md_blocks(md):
        text = _MD_INLINE.sub(r"\1", text)
        if kind.startswith("h"):
            lines += ["", text.upper() if kind == "h1" else text, ""]
        elif kind == "li":
            lines.append("• " + text)
        else:
            lines.append(text)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


def md_to_docx(path: str, md: str, title: str, note: str = "") -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=1)
    if note:
        p = doc.add_paragraph(note)
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x5B, 0x66, 0x80)

    def runs(par, text):
        pos = 0
        for m in _MD_INLINE.finditer(text):
            if m.start() > pos:
                par.add_run(text[pos:m.start()])
            par.add_run(m.group(1)).bold = True
            pos = m.end()
        if pos < len(text):
            par.add_run(text[pos:])

    for kind, text in _md_blocks(md):
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(_MD_INLINE.sub(r"\1", text), level=min(3, int(kind[1])))
        elif kind == "li":
            runs(doc.add_paragraph(style="List Bullet"), text)
        else:
            runs(doc.add_paragraph(), text)
    doc.core_properties.title = title
    doc.save(path)


def md_to_pdf(path: str, md: str, title: str, note: str = "") -> None:
    from fpdf import FPDF

    font = _find_font()
    if not font:
        raise RuntimeError("PDF için uygun yazı tipi bulunamadı.")
    regular, bold = font
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()
    pdf.add_font("Govde", "", regular)
    B = ""
    if bold:
        pdf.add_font("Govde", "B", bold)
        B = "B"
    pdf.set_font("Govde", B, 16)
    pdf.multi_cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    if note:
        pdf.set_font("Govde", "", 9)
        pdf.set_text_color(91, 102, 128)
        pdf.multi_cell(0, 5, note, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
    pdf.ln(3)
    md_ok = bool(bold)
    for kind, text in _md_blocks(md):
        if kind.startswith("h"):
            pdf.ln(2)
            pdf.set_font("Govde", B, 13 if kind != "h3" else 12)
            pdf.multi_cell(0, 7, _MD_INLINE.sub(r"\1", text), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif kind == "li":
            pdf.set_font("Govde", "", 11)
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(0, 6, "• " + (text if md_ok else _MD_INLINE.sub(r"\1", text)),
                           markdown=md_ok, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Govde", "", 11)
            pdf.multi_cell(0, 6, text if md_ok else _MD_INLINE.sub(r"\1", text),
                           markdown=md_ok, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)
    pdf.output(path)


def export_summary(path: str, md: str, title: str, note: str = "") -> None:
    p = path.lower()
    if p.endswith(".docx"):
        md_to_docx(path, md, title, note)
    elif p.endswith(".pdf"):
        md_to_pdf(path, md, title, note)
    elif p.endswith(".md"):
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"# {title}\n\n" + (f"_{note}_\n\n" if note else "") + md.strip() + "\n")
    else:
        with open(path, "w", encoding="utf-8") as f:
            f.write(title + "\n" + (note + "\n" if note else "") + "\n" + md_to_text(md))
